from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.core import mail
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from inventory import exports, services, statistiques
from inventory.models import Categorie, Fournisseur, Mouvement, Produit

# Backend mémoire : les emails sont collectés dans mail.outbox au lieu d'être envoyés.
PARAMETRES_EMAIL_TEST = {
    "EMAIL_BACKEND": "django.core.mail.backends.locmem.EmailBackend",
    "ALERTE_EMAIL_DESTINATAIRE": "gerant@stockflow.test",
    "DEFAULT_FROM_EMAIL": "stockflow@stockflow.test",
}

MOT_DE_PASSE_TEST = "motdepasse123"


def creer_utilisateur(nom, groupe=None):
    """Crée un utilisateur de test, éventuellement rattaché à un groupe métier."""
    utilisateur = User.objects.create_user(username=nom, password=MOT_DE_PASSE_TEST)
    if groupe:
        utilisateur.groups.add(Group.objects.get(name=groupe))
    return utilisateur


class BaseApplicationTestCase(TestCase):
    """
    Base des tests qui passent par l'interface.

    Les groupes Gerant/Magasinier sont créés par la commande de déploiement,
    pas par une migration : on l'exécute donc pour chaque classe de test.
    """

    @classmethod
    def setUpTestData(cls):
        call_command("initialiser_groupes", verbosity=0)


class EnregistrerMouvementTests(TestCase):
    """Tests de la logique métier de services.enregistrer_mouvement()."""

    def setUp(self):
        self.categorie = Categorie.objects.create(nom="Test")
        # On fixe directement quantite_stock ici : il s'agit de préparer l'état
        # initial d'un test, pas d'une opération métier.
        self.produit = Produit.objects.create(
            reference="TST-001",
            nom="Produit test",
            categorie=self.categorie,
            prix_achat=100,
            prix_vente=150,
            quantite_stock=10,
        )

    def test_entree_normale(self):
        """Stock 10 + entrée 5 = 15, et le mouvement est enregistré."""
        mouvement = services.enregistrer_mouvement(
            self.produit, Mouvement.ENTREE, 5, motif="Réception"
        )

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 15)
        self.assertEqual(Mouvement.objects.count(), 1)
        self.assertEqual(mouvement.type_mouvement, Mouvement.ENTREE)
        self.assertEqual(mouvement.quantite, 5)
        self.assertEqual(mouvement.produit, self.produit)

    def test_sortie_normale(self):
        """Stock 10 - sortie 4 = 6, et le mouvement est enregistré."""
        mouvement = services.enregistrer_mouvement(
            self.produit, Mouvement.SORTIE, 4, motif="Vente"
        )

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 6)
        self.assertEqual(Mouvement.objects.count(), 1)
        self.assertEqual(mouvement.type_mouvement, Mouvement.SORTIE)
        self.assertEqual(mouvement.quantite, 4)

    def test_sortie_superieure_au_stock(self):
        """Sortie de 15 sur un stock de 10 : refusée, stock inchangé."""
        with self.assertRaises(ValidationError):
            services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 15)

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 10)
        self.assertEqual(Mouvement.objects.count(), 0)

    def test_quantite_nulle_refusee(self):
        """Une quantité de 0 est refusée, en entrée comme en sortie."""
        for type_mouvement in (Mouvement.ENTREE, Mouvement.SORTIE):
            with self.subTest(type_mouvement=type_mouvement):
                with self.assertRaises(ValidationError):
                    services.enregistrer_mouvement(self.produit, type_mouvement, 0)

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 10)
        self.assertEqual(Mouvement.objects.count(), 0)

    def test_type_mouvement_inconnu_refuse(self):
        """Un type de mouvement hors ENTREE/SORTIE est refusé."""
        with self.assertRaises(ValidationError):
            services.enregistrer_mouvement(self.produit, "TRANSFERT", 5)

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 10)
        self.assertEqual(Mouvement.objects.count(), 0)


    def test_transaction_atomique_rollback_complet(self):
        """
        Si la création du Mouvement échoue, la mise à jour du stock est annulée.

        On simule une panne au moment de créer le mouvement et on vérifie qu'on
        ne se retrouve pas avec un stock modifié sans historique correspondant.
        """
        from unittest.mock import patch

        with patch.object(
            Mouvement.objects, "create", side_effect=RuntimeError("panne base de données")
        ):
            with self.assertRaises(RuntimeError):
                services.enregistrer_mouvement(self.produit, Mouvement.ENTREE, 5)

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 10)  # stock intact
        self.assertEqual(Mouvement.objects.count(), 0)


@override_settings(**PARAMETRES_EMAIL_TEST)
class AlerteStockTests(TestCase):
    """Déclenchement de l'alerte email au franchissement du seuil."""

    def setUp(self):
        self.categorie = Categorie.objects.create(nom="Test")
        self.produit = Produit.objects.create(
            reference="ALE-001",
            nom="Produit surveillé",
            categorie=self.categorie,
            prix_achat=100,
            prix_vente=150,
            quantite_stock=15,
            seuil_alerte=10,
        )

    def test_franchit_seuil_alerte(self):
        """La règle de franchissement, testée isolément."""
        self.assertTrue(services.franchit_seuil_alerte(15, 8, 10))    # franchissement
        self.assertTrue(services.franchit_seuil_alerte(11, 10, 10))   # atteint le seuil
        self.assertFalse(services.franchit_seuil_alerte(8, 6, 10))    # déjà sous le seuil
        self.assertFalse(services.franchit_seuil_alerte(20, 15, 10))  # reste au-dessus

    def test_sortie_qui_franchit_le_seuil_envoie_un_email(self):
        """Stock 15, seuil 10, sortie 7 -> stock 8 : une alerte part."""
        # captureOnCommitCallbacks exécute les callbacks transaction.on_commit,
        # qui ne se déclenchent pas seuls à l'intérieur d'un TestCase.
        with self.captureOnCommitCallbacks(execute=True):
            services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 7)

        self.assertEqual(len(mail.outbox), 1)
        message = mail.outbox[0]
        self.assertIn("Produit surveillé", message.subject)
        self.assertEqual(message.to, ["gerant@stockflow.test"])

        # L'email contient bien une version HTML en plus du texte.
        self.assertEqual(len(message.alternatives), 1)
        corps_html = message.alternatives[0][0]
        self.assertIn("text/html", message.alternatives[0][1])
        self.assertIn("Alerte de stock", corps_html)
        self.assertIn("STOCK CRITIQUE", corps_html)
        # Stock avant, quantité sortie et stock actuel sont présents.
        self.assertIn("15", corps_html)
        self.assertIn("7", corps_html)
        self.assertIn("8", corps_html)

    def test_produit_deja_sous_le_seuil_n_envoie_pas_d_email(self):
        """Stock 8 (déjà sous le seuil 10), sortie 2 -> aucune alerte."""
        self.produit.quantite_stock = 8
        self.produit.save(update_fields=["quantite_stock"])

        with self.captureOnCommitCallbacks(execute=True):
            services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 2)

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 6)
        self.assertEqual(len(mail.outbox), 0)

    def test_aucun_email_inutile(self):
        """Ni une entrée, ni une sortie qui reste au-dessus du seuil n'alertent."""
        with self.captureOnCommitCallbacks(execute=True):
            services.enregistrer_mouvement(self.produit, Mouvement.ENTREE, 50)
        self.assertEqual(len(mail.outbox), 0)

        with self.captureOnCommitCallbacks(execute=True):
            services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 5)
        self.assertEqual(len(mail.outbox), 0)

    def test_rupture_totale_signalee_comme_telle(self):
        """Stock tombé à 0 : l'email annonce une rupture, pas un stock critique."""
        with self.captureOnCommitCallbacks(execute=True):
            services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 15)

        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("RUPTURE", mail.outbox[0].alternatives[0][0])

    def test_sortie_refusee_n_envoie_aucun_email(self):
        """Un mouvement annulé ne doit générer aucune alerte."""
        with self.captureOnCommitCallbacks(execute=True):
            with self.assertRaises(ValidationError):
                services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 999)

        self.assertEqual(len(mail.outbox), 0)

    def test_panne_email_ne_casse_pas_la_sortie_de_stock(self):
        """Un SMTP en panne ne doit jamais faire échouer une vente."""
        from unittest.mock import patch

        with patch.object(
            services.EmailMultiAlternatives, "send", side_effect=OSError("SMTP injoignable")
        ):
            with self.captureOnCommitCallbacks(execute=True):
                services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 7)

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 8)  # la sortie est bien passée
        self.assertEqual(Mouvement.objects.count(), 1)


class AuthentificationTests(BaseApplicationTestCase):
    """Vérifie que les pages privées sont bien protégées."""

    def test_pages_privees_redirigent_vers_connexion(self):
        """AUCUNE page de l'application n'est accessible sans connexion."""
        categorie = Categorie.objects.create(nom="Test")
        produit = Produit.objects.create(
            reference="AUT-001", nom="Produit", categorie=categorie,
            prix_achat=100, prix_vente=150,
        )
        urls = [
            reverse("inventory:produit_liste"),
            reverse("inventory:produit_creer"),
            reverse("inventory:produit_detail", kwargs={"pk": produit.pk}),
            reverse("inventory:produit_modifier", kwargs={"pk": produit.pk}),
            reverse("inventory:mouvement_liste"),
            reverse("inventory:entree_stock"),
            reverse("inventory:sortie_stock"),
        ]
        for url in urls:
            with self.subTest(url=url):
                reponse = self.client.get(url)
                self.assertEqual(reponse.status_code, 302)
                self.assertIn(reverse("inventory:connexion"), reponse.url)

    def test_aucun_lien_prive_visible_sans_connexion(self):
        """La page de connexion ne doit exposer ni menu ni lien vers l'app."""
        reponse = self.client.get(reverse("inventory:connexion"))
        contenu = reponse.content.decode()

        self.assertNotIn(reverse("inventory:produit_liste"), contenu)
        self.assertNotIn(reverse("inventory:mouvement_liste"), contenu)
        self.assertNotIn(reverse("inventory:entree_stock"), contenu)
        self.assertNotIn(reverse("inventory:sortie_stock"), contenu)

    def test_utilisateur_connecte_ne_voit_pas_la_page_de_connexion(self):
        """Déjà connecté (y compris via /admin/) : /connexion/ renvoie vers l'app."""
        creer_utilisateur("testeur", groupe="Magasinier")
        self.client.login(username="testeur", password=MOT_DE_PASSE_TEST)

        reponse = self.client.get(reverse("inventory:connexion"))
        self.assertEqual(reponse.status_code, 302)
        self.assertEqual(reponse.url, reverse("inventory:produit_liste"))

    def test_page_connexion_accessible_sans_authentification(self):
        reponse = self.client.get(reverse("inventory:connexion"))
        self.assertEqual(reponse.status_code, 200)

    def test_connexion_par_le_formulaire(self):
        """Connexion via le vrai formulaire de la vue, pas via client.login()."""
        creer_utilisateur("testeur", groupe="Magasinier")

        reponse = self.client.post(
            reverse("inventory:connexion"),
            {"username": "testeur", "password": MOT_DE_PASSE_TEST},
        )
        self.assertEqual(reponse.status_code, 302)
        self.assertEqual(reponse.url, reverse("inventory:produit_liste"))

        # La page privée est maintenant accessible.
        reponse = self.client.get(reverse("inventory:produit_liste"))
        self.assertEqual(reponse.status_code, 200)

    def test_connexion_refusee_avec_mauvais_mot_de_passe(self):
        creer_utilisateur("testeur", groupe="Magasinier")

        reponse = self.client.post(
            reverse("inventory:connexion"),
            {"username": "testeur", "password": "mauvais"},
        )
        self.assertEqual(reponse.status_code, 200)
        self.assertContains(reponse, "Identifiants incorrects")


class ParcoursUtilisateurTests(BaseApplicationTestCase):
    """
    Parcours complet : connexion, création de produit, recherche, fiche,
    entrée, sortie, sortie refusée, historique, filtres, déconnexion.
    """

    def setUp(self):
        # Le parcours complet inclut la creation de produits : role Gerant.
        self.utilisateur = creer_utilisateur("testeur", groupe="Gerant")
        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.fournisseur = Fournisseur.objects.create(nom="Fournisseur test")

    def test_parcours_complet(self):
        # 1-2. Connexion
        self.assertTrue(self.client.login(username="testeur", password=MOT_DE_PASSE_TEST))

        # 3. Liste des produits
        reponse = self.client.get(reverse("inventory:produit_liste"))
        self.assertEqual(reponse.status_code, 200)

        # 4. Création d'un produit
        reponse = self.client.post(
            reverse("inventory:produit_creer"),
            {
                "reference": "PAR-001",
                "nom": "Riz parfumé 25kg",
                "categorie": self.categorie.pk,
                "fournisseur": self.fournisseur.pk,
                "prix_achat": "12000",
                "prix_vente": "14000",
                "seuil_alerte": "5",
                "actif": "on",
            },
        )
        self.assertEqual(reponse.status_code, 302)
        produit = Produit.objects.get(reference="PAR-001")
        self.assertEqual(produit.quantite_stock, 0)

        # 5-6. Le produit est dans la liste et retrouvable par recherche
        reponse = self.client.get(reverse("inventory:produit_liste"), {"q": "Riz parfumé"})
        self.assertContains(reponse, "PAR-001")
        reponse = self.client.get(reverse("inventory:produit_liste"), {"q": "PAR-001"})
        self.assertContains(reponse, "Riz parfumé")

        # 7. Fiche produit
        reponse = self.client.get(produit.get_absolute_url())
        self.assertEqual(reponse.status_code, 200)

        # 8-9. Entrée de stock : le stock augmente
        reponse = self.client.post(
            reverse("inventory:entree_stock"),
            {"produit": produit.pk, "quantite": "20", "motif": "Réception"},
        )
        self.assertEqual(reponse.status_code, 302)
        produit.refresh_from_db()
        self.assertEqual(produit.quantite_stock, 20)

        # 10-11. Sortie normale : le stock diminue
        reponse = self.client.post(
            reverse("inventory:sortie_stock"),
            {"produit": produit.pk, "quantite": "8", "motif": "Vente"},
        )
        self.assertEqual(reponse.status_code, 302)
        produit.refresh_from_db()
        self.assertEqual(produit.quantite_stock, 12)

        # 12-14. Sortie supérieure au stock : refusée proprement, stock inchangé
        reponse = self.client.post(
            reverse("inventory:sortie_stock"),
            {"produit": produit.pk, "quantite": "999", "motif": "Vente"},
        )
        self.assertEqual(reponse.status_code, 200)  # formulaire réaffiché
        self.assertContains(reponse, "Stock insuffisant")
        produit.refresh_from_db()
        self.assertEqual(produit.quantite_stock, 12)
        self.assertEqual(Mouvement.objects.count(), 2)  # seuls l'entrée et la sortie valide

        # 15. Historique
        reponse = self.client.get(reverse("inventory:mouvement_liste"))
        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(len(reponse.context["mouvements"]), 2)

        # 16. Filtres combinés de l'historique
        reponse = self.client.get(
            reverse("inventory:mouvement_liste"),
            {"produit": produit.pk, "type": Mouvement.SORTIE},
        )
        self.assertEqual(len(reponse.context["mouvements"]), 1)
        self.assertEqual(reponse.context["mouvements"][0].type_mouvement, Mouvement.SORTIE)

        # 17. Les 20 derniers mouvements sur la fiche produit
        reponse = self.client.get(produit.get_absolute_url())
        self.assertEqual(len(reponse.context["mouvements"]), 2)

        # 18. Déconnexion (POST obligatoire depuis Django 5)
        reponse = self.client.post(reverse("inventory:deconnexion"))
        self.assertEqual(reponse.status_code, 302)
        reponse = self.client.get(reverse("inventory:produit_liste"))
        self.assertEqual(reponse.status_code, 302)  # redirigé vers la connexion

    def test_fiche_produit_limite_a_20_mouvements(self):
        self.client.login(username="testeur", password=MOT_DE_PASSE_TEST)
        produit = Produit.objects.create(
            reference="LIM-001",
            nom="Produit limite",
            categorie=self.categorie,
            prix_achat=100,
            prix_vente=150,
        )
        for _ in range(25):
            services.enregistrer_mouvement(produit, Mouvement.ENTREE, 1)

        reponse = self.client.get(produit.get_absolute_url())
        self.assertEqual(len(reponse.context["mouvements"]), 20)


class ProduitListeFiltresTests(BaseApplicationTestCase):
    """Filtres et recherche de la liste des produits."""

    def setUp(self):
        self.utilisateur = creer_utilisateur("testeur", groupe="Magasinier")
        self.client.login(username="testeur", password=MOT_DE_PASSE_TEST)
        self.categorie_a = Categorie.objects.create(nom="Alimentation")
        self.categorie_b = Categorie.objects.create(nom="Boissons")
        self.produit_a = Produit.objects.create(
            reference="A-001", nom="Riz", categorie=self.categorie_a,
            prix_achat=100, prix_vente=150, quantite_stock=100, seuil_alerte=10,
        )
        self.produit_b = Produit.objects.create(
            reference="B-001", nom="Jus", categorie=self.categorie_b,
            prix_achat=100, prix_vente=150, quantite_stock=2, seuil_alerte=10, actif=False,
        )

    def test_filtre_categorie(self):
        reponse = self.client.get(reverse("inventory:produit_liste"), {"categorie": self.categorie_b.pk})
        self.assertEqual(list(reponse.context["produits"]), [self.produit_b])

    def test_filtre_statut_actif(self):
        reponse = self.client.get(reverse("inventory:produit_liste"), {"actif": "1"})
        self.assertEqual(list(reponse.context["produits"]), [self.produit_a])

    def test_filtre_stock_bas(self):
        reponse = self.client.get(reverse("inventory:produit_liste"), {"stock_bas": "1"})
        self.assertEqual(list(reponse.context["produits"]), [self.produit_b])

    def test_filtre_invalide_est_ignore_sans_erreur(self):
        reponse = self.client.get(reverse("inventory:produit_liste"), {"categorie": "abc"})
        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(len(reponse.context["produits"]), 2)


class TableauBordTests(BaseApplicationTestCase):
    """KPI, produits en alerte et derniers mouvements du tableau de bord."""

    def setUp(self):
        self.gerant = creer_utilisateur("gerant", groupe="Gerant")
        self.client.login(username="gerant", password=MOT_DE_PASSE_TEST)

        self.categorie = Categorie.objects.create(nom="Alimentation")
        # 10 x 1000 = 10 000 de valeur, stock normal
        self.produit_normal = Produit.objects.create(
            reference="TB-001", nom="Produit normal", categorie=self.categorie,
            prix_achat=Decimal("1000.00"), prix_vente=Decimal("1500.00"),
            quantite_stock=10, seuil_alerte=5,
        )
        # 3 <= 5 : sous seuil. 3 x 500 = 1 500
        self.produit_sous_seuil = Produit.objects.create(
            reference="TB-002", nom="Produit sous seuil", categorie=self.categorie,
            prix_achat=Decimal("500.00"), prix_vente=Decimal("800.00"),
            quantite_stock=3, seuil_alerte=5,
        )
        # 0 : rupture. Valeur 0
        self.produit_rupture = Produit.objects.create(
            reference="TB-003", nom="Produit en rupture", categorie=self.categorie,
            prix_achat=Decimal("200.00"), prix_vente=Decimal("300.00"),
            quantite_stock=0, seuil_alerte=5,
        )
        # Inactif : ne doit compter dans aucun KPI
        self.produit_inactif = Produit.objects.create(
            reference="TB-004", nom="Produit inactif", categorie=self.categorie,
            prix_achat=Decimal("9999.00"), prix_vente=Decimal("9999.00"),
            quantite_stock=50, seuil_alerte=5, actif=False,
        )

    def test_calcul_des_kpis(self):
        kpis = statistiques.kpis_stock()

        self.assertEqual(kpis["produits_actifs"], 3)                 # l'inactif est exclu
        self.assertEqual(kpis["quantite_totale"], 13)                # 10 + 3 + 0
        self.assertEqual(kpis["valeur_stock"], Decimal("11500.00"))  # 10000 + 1500 + 0
        self.assertEqual(kpis["produits_en_alerte"], 2)              # sous seuil + rupture
        self.assertEqual(kpis["produits_en_rupture"], 1)

    def test_kpis_sur_base_vide(self):
        """Aucun produit : les agrégats valent 0, pas None."""
        Produit.objects.all().delete()
        kpis = statistiques.kpis_stock()

        self.assertEqual(kpis["produits_actifs"], 0)
        self.assertEqual(kpis["quantite_totale"], 0)
        self.assertEqual(kpis["valeur_stock"], Decimal("0.00"))

    def test_produits_en_alerte(self):
        alertes = list(statistiques.produits_en_alerte())

        # Triés par stock croissant : la rupture d'abord.
        self.assertEqual(alertes, [self.produit_rupture, self.produit_sous_seuil])
        self.assertNotIn(self.produit_normal, alertes)
        self.assertNotIn(self.produit_inactif, alertes)

    def test_derniers_mouvements_limite_et_ordre(self):
        for i in range(12):
            services.enregistrer_mouvement(self.produit_normal, Mouvement.ENTREE, i + 1)

        derniers = list(statistiques.derniers_mouvements(10))

        self.assertEqual(len(derniers), 10)
        self.assertEqual(derniers[0].quantite, 12)  # le plus récent en premier

    def test_page_tableau_bord_affiche_kpis_et_alertes(self):
        reponse = self.client.get(reverse("inventory:tableau_bord"))

        self.assertEqual(reponse.status_code, 200)
        self.assertEqual(reponse.context["kpis"]["produits_actifs"], 3)
        self.assertEqual(len(reponse.context["produits_en_alerte"]), 2)
        self.assertContains(reponse, "Produit en rupture")
        self.assertContains(reponse, "Rupture")

    def test_etat_vide_aucune_alerte(self):
        """Sans produit en alerte, un message remplace le tableau."""
        self.produit_sous_seuil.delete()
        self.produit_rupture.delete()

        reponse = self.client.get(reverse("inventory:tableau_bord"))
        self.assertContains(reponse, "Aucun produit en alerte")

    def test_etat_vide_aucun_mouvement(self):
        reponse = self.client.get(reverse("inventory:tableau_bord"))
        self.assertContains(reponse, "Aucun mouvement enregistr")


class ExportExcelTests(BaseApplicationTestCase):
    """Génération et formatage du fichier Excel."""

    def setUp(self):
        self.gerant = creer_utilisateur("gerant", groupe="Gerant")
        self.client.login(username="gerant", password=MOT_DE_PASSE_TEST)

        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.fournisseur = Fournisseur.objects.create(nom="Fournisseur A")
        self.produit_normal = Produit.objects.create(
            reference="EX-001", nom="Riz 25kg", categorie=self.categorie,
            fournisseur=self.fournisseur, prix_achat=Decimal("12000.00"),
            prix_vente=Decimal("14000.00"), quantite_stock=10, seuil_alerte=5,
        )
        self.produit_rupture = Produit.objects.create(
            reference="EX-002", nom="Huile 5L", categorie=self.categorie,
            prix_achat=Decimal("5000.00"), prix_vente=Decimal("6000.00"),
            quantite_stock=0, seuil_alerte=5,
        )

    def _telecharger_classeur(self, parametres=None):
        reponse = self.client.get(reverse("inventory:export_stock_excel"), parametres or {})
        self.assertEqual(reponse.status_code, 200)
        return reponse, load_workbook(BytesIO(reponse.content))

    def test_generation_du_fichier(self):
        reponse, classeur = self._telecharger_classeur()

        self.assertEqual(
            reponse["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", reponse["Content-Disposition"])
        self.assertIn(".xlsx", reponse["Content-Disposition"])
        self.assertEqual(classeur.active.title, "État du stock")

    def test_presence_des_colonnes(self):
        _, classeur = self._telecharger_classeur()
        feuille = classeur.active

        entetes = [cellule.value for cellule in feuille[4]]
        self.assertEqual(entetes, [titre for titre, _ in exports.COLONNES])

    def test_donnees_et_statuts(self):
        _, classeur = self._telecharger_classeur()
        feuille = classeur.active

        lignes = {feuille.cell(row=l, column=1).value: l for l in range(5, feuille.max_row + 1)}
        self.assertEqual(set(lignes), {"EX-001", "EX-002"})

        ligne_normale = lignes["EX-001"]
        self.assertEqual(feuille.cell(row=ligne_normale, column=2).value, "Riz 25kg")
        self.assertEqual(feuille.cell(row=ligne_normale, column=3).value, "Alimentation")
        self.assertEqual(feuille.cell(row=ligne_normale, column=4).value, "Fournisseur A")
        self.assertEqual(feuille.cell(row=ligne_normale, column=5).value, 10)
        self.assertEqual(feuille.cell(row=ligne_normale, column=7).value, "Stock normal")
        # Valeur du stock = quantité x prix d'achat
        self.assertEqual(feuille.cell(row=ligne_normale, column=9).value, Decimal("120000.00"))

        self.assertEqual(feuille.cell(row=lignes["EX-002"], column=7).value, "Rupture")

    def test_formatage_de_base(self):
        _, classeur = self._telecharger_classeur()
        feuille = classeur.active

        # En-têtes en gras
        self.assertTrue(feuille.cell(row=4, column=1).font.bold)
        # Filtre automatique et volets figés
        self.assertIsNotNone(feuille.auto_filter.ref)
        self.assertEqual(feuille.freeze_panes, "A5")
        # Largeur de colonne adaptée
        self.assertEqual(feuille.column_dimensions["B"].width, 34)
        # Format monétaire sur la valeur du stock
        self.assertEqual(feuille.cell(row=5, column=9).number_format, "#,##0.00")
        # Statut coloré
        self.assertNotEqual(feuille.cell(row=5, column=7).fill.fgColor.rgb, "00000000")
        # Date de génération présente
        self.assertIn("Document", feuille["A2"].value)

    def test_export_respecte_les_filtres(self):
        """L'export porte sur ce que l'utilisateur a filtré à l'écran."""
        _, classeur = self._telecharger_classeur({"stock_bas": "1"})
        feuille = classeur.active

        references = [
            feuille.cell(row=l, column=1).value for l in range(5, feuille.max_row + 1)
        ]
        self.assertEqual(references, ["EX-002"])  # seul le produit en rupture


class PermissionsTests(BaseApplicationTestCase):
    """Les restrictions sont appliquées côté serveur, pas seulement masquées."""

    def setUp(self):
        creer_utilisateur("gerant", groupe="Gerant")
        creer_utilisateur("magasinier", groupe="Magasinier")
        creer_utilisateur("sans_role")  # aucun groupe

        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.produit = Produit.objects.create(
            reference="PER-001", nom="Produit", categorie=self.categorie,
            prix_achat=100, prix_vente=150, quantite_stock=10,
        )

    def _connecter(self, nom):
        self.assertTrue(self.client.login(username=nom, password=MOT_DE_PASSE_TEST))

    def test_acces_gerant(self):
        """Le gérant accède à tout, y compris aux fonctions sensibles."""
        self._connecter("gerant")
        urls = [
            reverse("inventory:tableau_bord"),
            reverse("inventory:export_stock_excel"),
            reverse("inventory:produit_liste"),
            reverse("inventory:produit_creer"),
            reverse("inventory:produit_modifier", kwargs={"pk": self.produit.pk}),
            reverse("inventory:mouvement_liste"),
            reverse("inventory:entree_stock"),
            reverse("inventory:sortie_stock"),
        ]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_acces_magasinier(self):
        """Le magasinier gère le quotidien : consultation et mouvements."""
        self._connecter("magasinier")
        urls_autorisees = [
            reverse("inventory:produit_liste"),
            reverse("inventory:produit_detail", kwargs={"pk": self.produit.pk}),
            reverse("inventory:mouvement_liste"),
            reverse("inventory:entree_stock"),
            reverse("inventory:sortie_stock"),
        ]
        for url in urls_autorisees:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_acces_refuse_au_magasinier_sur_les_vues_sensibles(self):
        """Tableau de bord, export et gestion des produits lui sont interdits."""
        self._connecter("magasinier")
        urls_interdites = [
            reverse("inventory:tableau_bord"),
            reverse("inventory:export_stock_excel"),
            reverse("inventory:produit_creer"),
            reverse("inventory:produit_modifier", kwargs={"pk": self.produit.pk}),
        ]
        for url in urls_interdites:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 403)

    def test_magasinier_ne_peut_pas_creer_un_produit_en_post(self):
        """La restriction tient aussi en POST, pas seulement à l'affichage."""
        self._connecter("magasinier")
        reponse = self.client.post(
            reverse("inventory:produit_creer"),
            {
                "reference": "HACK-001", "nom": "Produit interdit",
                "categorie": self.categorie.pk, "prix_achat": "1", "prix_vente": "2",
                "seuil_alerte": "1", "actif": "on",
            },
        )
        self.assertEqual(reponse.status_code, 403)
        self.assertFalse(Produit.objects.filter(reference="HACK-001").exists())

    def test_utilisateur_sans_role_na_acces_a_rien(self):
        self._connecter("sans_role")
        for url in [
            reverse("inventory:produit_liste"),
            reverse("inventory:mouvement_liste"),
            reverse("inventory:tableau_bord"),
            reverse("inventory:export_stock_excel"),
        ]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 403)

    def test_liens_sensibles_masques_pour_le_magasinier(self):
        """L'interface est cohérente avec les droits (en plus du contrôle serveur)."""
        self._connecter("magasinier")
        contenu = self.client.get(reverse("inventory:produit_liste")).content.decode()

        self.assertNotIn(reverse("inventory:tableau_bord"), contenu)
        self.assertNotIn(reverse("inventory:export_stock_excel"), contenu)
        self.assertNotIn(reverse("inventory:produit_creer"), contenu)


@override_settings(**PARAMETRES_EMAIL_TEST)
class RapportQuotidienTests(BaseApplicationTestCase):
    """Commande rapport_quotidien."""

    def setUp(self):
        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.produit = Produit.objects.create(
            reference="RAP-001", nom="Riz 25kg", categorie=self.categorie,
            prix_achat=Decimal("12000.00"), prix_vente=Decimal("14000.00"),
            quantite_stock=20, seuil_alerte=5,
        )

    def test_statistiques_du_jour(self):
        services.enregistrer_mouvement(self.produit, Mouvement.ENTREE, 10)
        services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 4)
        services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 6)

        stats = statistiques.statistiques_du_jour()

        self.assertEqual(stats["entrees_nombre"], 1)
        self.assertEqual(stats["entrees_quantite"], 10)
        self.assertEqual(stats["sorties_nombre"], 2)
        self.assertEqual(stats["sorties_quantite"], 10)
        self.assertEqual(stats["mouvements_nombre"], 3)

    def test_commande_envoie_le_rapport(self):
        services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 18)  # passe sous le seuil
        mail.outbox = []

        call_command("rapport_quotidien", verbosity=0)

        self.assertEqual(len(mail.outbox), 1)
        message = mail.outbox[0]
        self.assertIn("Rapport quotidien", message.subject)
        self.assertEqual(message.to, ["gerant@stockflow.test"])

        corps_html = message.alternatives[0][0]
        self.assertIn("Rapport quotidien", corps_html)
        self.assertIn("Riz 25kg", corps_html)  # produit en alerte listé
        self.assertIn("Mouvements de la journ", corps_html)

        # La version texte existe aussi.
        self.assertIn("RAPPORT QUOTIDIEN", message.body)

    def test_commande_mode_apercu_n_envoie_rien(self):
        call_command("rapport_quotidien", "--apercu", verbosity=0)
        self.assertEqual(len(mail.outbox), 0)

    def test_commande_refuse_une_date_invalide(self):
        from django.core.management.base import CommandError

        with self.assertRaises(CommandError):
            call_command("rapport_quotidien", "--jour", "pas-une-date", verbosity=0)

    @override_settings(ALERTE_EMAIL_DESTINATAIRE="")
    def test_commande_echoue_proprement_sans_destinataire(self):
        from django.core.management.base import CommandError

        with self.assertRaises(CommandError):
            call_command("rapport_quotidien", verbosity=0)


class InitialiserGroupesTests(TestCase):
    """La commande de création des groupes est rejouable sans effet de bord."""

    def test_commande_idempotente(self):
        call_command("initialiser_groupes", verbosity=0)
        call_command("initialiser_groupes", verbosity=0)

        self.assertEqual(Group.objects.filter(name__in=["Gerant", "Magasinier"]).count(), 2)
        gerant = Group.objects.get(name="Gerant")
        magasinier = Group.objects.get(name="Magasinier")

        self.assertEqual(gerant.permissions.count(), 13)
        self.assertEqual(magasinier.permissions.count(), 5)

        codes_magasinier = set(magasinier.permissions.values_list("codename", flat=True))
        self.assertNotIn("acceder_tableau_bord", codes_magasinier)
        self.assertNotIn("exporter_stock", codes_magasinier)
        self.assertNotIn("add_produit", codes_magasinier)


@override_settings(**PARAMETRES_EMAIL_TEST)
class EncodageEmailsTests(TestCase):
    """
    Non-regression : le backend console de Django ecrit sur la sortie standard,
    encodee en cp1252 sous Windows. Un caractere hors cp1252 dans un template
    email (ex : le signe moins typographique U+2212) faisait echouer l'envoi
    silencieusement en developpement.
    """

    def setUp(self):
        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.produit = Produit.objects.create(
            reference="ENC-001", nom="Produit accentue", categorie=self.categorie,
            prix_achat=Decimal("1000.00"), prix_vente=Decimal("1500.00"),
            quantite_stock=15, seuil_alerte=10,
        )

    def _verifier_encodable(self, message):
        for contenu in [message.subject, message.body] + [a[0] for a in message.alternatives]:
            contenu.encode("cp1252")  # leve UnicodeEncodeError si un caractere passe

    def test_email_alerte_encodable_en_console_windows(self):
        with self.captureOnCommitCallbacks(execute=True):
            services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 7)

        self.assertEqual(len(mail.outbox), 1)
        self._verifier_encodable(mail.outbox[0])

    def test_email_rapport_encodable_en_console_windows(self):
        call_command("rapport_quotidien", verbosity=0)

        self.assertEqual(len(mail.outbox), 1)
        self._verifier_encodable(mail.outbox[0])


class TemplatesTests(BaseApplicationTestCase):
    """
    Non-regression : en Django, {# ... #} ne commente QUE sur une seule ligne.
    Un commentaire etale sur plusieurs lignes n'est pas reconnu et s'affiche
    tel quel dans la page (bug constate sur le menu de navigation).
    """

    def setUp(self):
        creer_utilisateur("gerant", groupe="Gerant")
        self.client.login(username="gerant", password=MOT_DE_PASSE_TEST)
        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.produit = Produit.objects.create(
            reference="TPL-001", nom="Produit", categorie=self.categorie,
            prix_achat=100, prix_vente=150, quantite_stock=10,
        )

    def test_aucun_marqueur_de_template_dans_les_pages(self):
        urls = [
            reverse("inventory:tableau_bord"),
            reverse("inventory:produit_liste"),
            reverse("inventory:produit_detail", kwargs={"pk": self.produit.pk}),
            reverse("inventory:produit_creer"),
            reverse("inventory:mouvement_liste"),
            reverse("inventory:entree_stock"),
            reverse("inventory:sortie_stock"),
        ]
        for url in urls:
            with self.subTest(url=url):
                contenu = self.client.get(url).content.decode()
                for marqueur in ("{#", "#}", "{%", "%}", "{{", "}}"):
                    self.assertNotIn(
                        marqueur, contenu,
                        f"Marqueur de template {marqueur!r} visible dans la page {url}",
                    )

    def test_page_connexion_sans_marqueur_de_template(self):
        self.client.logout()
        contenu = self.client.get(reverse("inventory:connexion")).content.decode()
        for marqueur in ("{#", "#}", "{%", "%}", "{{", "}}"):
            self.assertNotIn(marqueur, contenu)


class ValeurParCategorieTests(BaseApplicationTestCase):
    """Ventilation de la valeur du stock par categorie (tableau de bord)."""

    def setUp(self):
        self.gerant = creer_utilisateur("gerant", groupe="Gerant")
        self.client.login(username="gerant", password=MOT_DE_PASSE_TEST)

        self.alimentation = Categorie.objects.create(nom="Alimentation")
        self.boissons = Categorie.objects.create(nom="Boissons")
        Categorie.objects.create(nom="Categorie vide")

        # 10 x 1000 = 10 000
        Produit.objects.create(
            reference="VC-001", nom="Riz", categorie=self.alimentation,
            prix_achat=Decimal("1000.00"), prix_vente=Decimal("1500.00"), quantite_stock=10,
        )
        # 4 x 500 = 2 000
        Produit.objects.create(
            reference="VC-002", nom="Jus", categorie=self.boissons,
            prix_achat=Decimal("500.00"), prix_vente=Decimal("800.00"), quantite_stock=4,
        )
        # Produit desactive : ne doit pas compter
        Produit.objects.create(
            reference="VC-003", nom="Ancien", categorie=self.boissons,
            prix_achat=Decimal("9999.00"), prix_vente=Decimal("9999.00"),
            quantite_stock=50, actif=False,
        )

    def test_valeurs_et_ordre(self):
        categories = statistiques.valeur_par_categorie()

        self.assertEqual([c.nom for c in categories], ["Alimentation", "Boissons"])
        self.assertEqual(categories[0].valeur, Decimal("10000.00"))
        self.assertEqual(categories[1].valeur, Decimal("2000.00"))

    def test_categorie_sans_stock_exclue(self):
        noms = [c.nom for c in statistiques.valeur_par_categorie()]
        self.assertNotIn("Categorie vide", noms)

    def test_pourcentages_relatifs_au_maximum(self):
        categories = statistiques.valeur_par_categorie()

        self.assertEqual(categories[0].pourcentage, 100)
        self.assertEqual(categories[1].pourcentage, 20)  # 2000 / 10000

    def test_affichage_sur_le_tableau_de_bord(self):
        reponse = self.client.get(reverse("inventory:tableau_bord"))

        self.assertEqual(reponse.status_code, 200)
        self.assertContains(reponse, "Valeur du stock par catégorie")
        self.assertContains(reponse, "Alimentation")

    def test_etat_vide_sans_aucun_stock(self):
        Produit.objects.all().update(quantite_stock=0)

        reponse = self.client.get(reverse("inventory:tableau_bord"))
        self.assertContains(reponse, "Aucune valeur à afficher")


class TracabiliteMouvementTests(BaseApplicationTestCase):
    """Champs de traçabilité : utilisateur, stock après, document, destination."""

    def setUp(self):
        self.magasinier = creer_utilisateur("magasinier", groupe="Magasinier")
        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.produit = Produit.objects.create(
            reference="TRA-001", nom="Riz 25kg", categorie=self.categorie,
            prix_achat=Decimal("1000.00"), prix_vente=Decimal("1400.00"),
            quantite_stock=50, seuil_alerte=10,
        )

    def test_service_enregistre_l_utilisateur(self):
        mouvement = services.enregistrer_mouvement(
            self.produit, Mouvement.SORTIE, 5, motif="Vente",
            utilisateur=self.magasinier,
        )
        self.assertEqual(mouvement.utilisateur, self.magasinier)

    def test_utilisateur_facultatif(self):
        """Les scripts (seed, import) peuvent enregistrer sans utilisateur."""
        mouvement = services.enregistrer_mouvement(self.produit, Mouvement.ENTREE, 5)
        self.assertIsNone(mouvement.utilisateur)

    def test_stock_apres_est_fige(self):
        """stock_apres garde la valeur du stock au moment du mouvement."""
        premier = services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 20)
        second = services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 10)

        self.assertEqual(premier.stock_apres, 30)
        self.assertEqual(second.stock_apres, 20)

        # Le stock du produit continue d'évoluer, l'historique ne bouge plus.
        services.enregistrer_mouvement(self.produit, Mouvement.ENTREE, 100)
        premier.refresh_from_db()
        self.assertEqual(premier.stock_apres, 30)

    def test_document_et_destination_enregistres(self):
        mouvement = services.enregistrer_mouvement(
            self.produit, Mouvement.SORTIE, 5, motif="Vente",
            document="BON-0412", destination="Chantier Godomey",
        )
        self.assertEqual(mouvement.document, "BON-0412")
        self.assertEqual(mouvement.destination, "Chantier Godomey")

    def test_vue_sortie_enregistre_l_utilisateur_connecte(self):
        """La traçabilité vient de la session, pas d'un champ du formulaire."""
        self.client.login(username="magasinier", password=MOT_DE_PASSE_TEST)

        reponse = self.client.post(
            reverse("inventory:sortie_stock"),
            {
                "produit": self.produit.pk, "quantite": "5", "motif": "Vente",
                "document": "BON-0001", "destination": "Client Sossou",
            },
        )
        self.assertEqual(reponse.status_code, 302)

        mouvement = Mouvement.objects.get()
        self.assertEqual(mouvement.utilisateur, self.magasinier)
        self.assertEqual(mouvement.stock_apres, 45)
        self.assertEqual(mouvement.document, "BON-0001")
        self.assertEqual(mouvement.destination, "Client Sossou")

    def test_colonnes_visibles_dans_l_historique(self):
        self.client.login(username="magasinier", password=MOT_DE_PASSE_TEST)
        services.enregistrer_mouvement(
            self.produit, Mouvement.SORTIE, 5, motif="Vente",
            utilisateur=self.magasinier, document="BON-0412",
            destination="Chantier Godomey",
        )

        reponse = self.client.get(reverse("inventory:mouvement_liste"))

        self.assertContains(reponse, "BON-0412")
        self.assertContains(reponse, "Chantier Godomey")
        self.assertContains(reponse, "magasinier")
        self.assertContains(reponse, "Stock après")

    def test_mouvement_ancien_sans_stock_apres_s_affiche(self):
        """Les mouvements antérieurs au champ ne cassent pas l'affichage."""
        self.client.login(username="magasinier", password=MOT_DE_PASSE_TEST)
        mouvement = services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 5)
        Mouvement.objects.filter(pk=mouvement.pk).update(stock_apres=None)

        reponse = self.client.get(reverse("inventory:mouvement_liste"))
        self.assertEqual(reponse.status_code, 200)


class AjustementStockTests(BaseApplicationTestCase):
    """Ajustement : constat de perte, casse ou écart d'inventaire."""

    def setUp(self):
        self.magasinier = creer_utilisateur("magasinier", groupe="Magasinier")
        self.client.login(username="magasinier", password=MOT_DE_PASSE_TEST)
        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.produit = Produit.objects.create(
            reference="AJU-001", nom="Bouteille", categorie=self.categorie,
            prix_achat=Decimal("500.00"), prix_vente=Decimal("800.00"),
            quantite_stock=50, seuil_alerte=10,
        )

    def test_ajustement_diminue_le_stock(self):
        mouvement = services.enregistrer_mouvement(
            self.produit, Mouvement.AJUSTEMENT, 3, motif="Casse",
        )

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 47)
        self.assertEqual(mouvement.type_mouvement, Mouvement.AJUSTEMENT)
        self.assertEqual(mouvement.stock_apres, 47)

    def test_motif_obligatoire(self):
        """Un ajustement sans justification est refusé."""
        with self.assertRaises(ValidationError):
            services.enregistrer_mouvement(self.produit, Mouvement.AJUSTEMENT, 3)
        with self.assertRaises(ValidationError):
            services.enregistrer_mouvement(self.produit, Mouvement.AJUSTEMENT, 3, motif="   ")

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 50)
        self.assertEqual(Mouvement.objects.count(), 0)

    def test_ajustement_superieur_au_stock_refuse(self):
        with self.assertRaises(ValidationError):
            services.enregistrer_mouvement(
                self.produit, Mouvement.AJUSTEMENT, 999, motif="Casse",
            )

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 50)

    @override_settings(**PARAMETRES_EMAIL_TEST)
    def test_ajustement_qui_franchit_le_seuil_alerte(self):
        """Une casse qui fait passer sous le seuil prévient le gérant."""
        with self.captureOnCommitCallbacks(execute=True):
            services.enregistrer_mouvement(
                self.produit, Mouvement.AJUSTEMENT, 45, motif="Casse",
            )

        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 5)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Bouteille", mail.outbox[0].subject)

    def test_formulaire_ajustement_exige_un_motif(self):
        reponse = self.client.post(
            reverse("inventory:ajustement_stock"),
            {"produit": self.produit.pk, "quantite": "3", "motif": ""},
        )

        self.assertEqual(reponse.status_code, 200)  # formulaire réaffiché
        self.produit.refresh_from_db()
        self.assertEqual(self.produit.quantite_stock, 50)
        self.assertEqual(Mouvement.objects.count(), 0)

    def test_formulaire_ajustement_sans_champ_destination(self):
        """Un ajustement n'a ni client ni destination."""
        reponse = self.client.get(reverse("inventory:ajustement_stock"))
        self.assertNotIn("destination", reponse.context["form"].fields)

    def test_vue_ajustement_complete(self):
        reponse = self.client.post(
            reverse("inventory:ajustement_stock"),
            {"produit": self.produit.pk, "quantite": "3", "motif": "Casse au dépôt"},
        )

        self.assertEqual(reponse.status_code, 302)
        mouvement = Mouvement.objects.get()
        self.assertEqual(mouvement.type_mouvement, Mouvement.AJUSTEMENT)
        self.assertEqual(mouvement.motif, "Casse au dépôt")
        self.assertEqual(mouvement.utilisateur, self.magasinier)

    def test_badge_ajustement_dans_l_historique(self):
        services.enregistrer_mouvement(
            self.produit, Mouvement.AJUSTEMENT, 3, motif="Casse",
        )
        reponse = self.client.get(reverse("inventory:mouvement_liste"))
        self.assertContains(reponse, "Ajustement")


class FiltresHistoriqueTests(BaseApplicationTestCase):
    """Recherche libre, filtre utilisateur et filtre par type dans l'historique."""

    def setUp(self):
        self.gerant = creer_utilisateur("gerant", groupe="Gerant")
        self.magasinier = creer_utilisateur("magasinier", groupe="Magasinier")
        self.client.login(username="gerant", password=MOT_DE_PASSE_TEST)

        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.riz = Produit.objects.create(
            reference="FIL-RIZ", nom="Riz 25kg", categorie=self.categorie,
            prix_achat=Decimal("1000.00"), prix_vente=Decimal("1400.00"), quantite_stock=100,
        )
        self.huile = Produit.objects.create(
            reference="FIL-HUI", nom="Huile 5L", categorie=self.categorie,
            prix_achat=Decimal("500.00"), prix_vente=Decimal("700.00"), quantite_stock=100,
        )

        services.enregistrer_mouvement(
            self.riz, Mouvement.SORTIE, 5, motif="Vente",
            utilisateur=self.gerant, document="BON-1000",
        )
        services.enregistrer_mouvement(
            self.huile, Mouvement.ENTREE, 10, motif="Livraison",
            utilisateur=self.magasinier, document="BL-2000",
        )
        services.enregistrer_mouvement(
            self.huile, Mouvement.AJUSTEMENT, 2, motif="Casse",
            utilisateur=self.magasinier,
        )

    def _mouvements(self, parametres):
        reponse = self.client.get(reverse("inventory:mouvement_liste"), parametres)
        self.assertEqual(reponse.status_code, 200)
        return list(reponse.context["mouvements"])

    def test_filtre_par_utilisateur(self):
        resultats = self._mouvements({"utilisateur": self.magasinier.pk})

        self.assertEqual(len(resultats), 2)
        self.assertTrue(all(m.utilisateur == self.magasinier for m in resultats))

    def test_recherche_par_reference_de_bon(self):
        resultats = self._mouvements({"q": "BON-1000"})

        self.assertEqual(len(resultats), 1)
        self.assertEqual(resultats[0].document, "BON-1000")

    def test_recherche_par_nom_de_produit(self):
        resultats = self._mouvements({"q": "Huile"})
        self.assertEqual(len(resultats), 2)

    def test_recherche_par_reference_produit(self):
        resultats = self._mouvements({"q": "FIL-RIZ"})
        self.assertEqual(len(resultats), 1)

    def test_filtre_par_type_ajustement(self):
        resultats = self._mouvements({"type": Mouvement.AJUSTEMENT})

        self.assertEqual(len(resultats), 1)
        self.assertEqual(resultats[0].motif, "Casse")

    def test_filtres_combines(self):
        resultats = self._mouvements({
            "utilisateur": self.magasinier.pk,
            "type": Mouvement.ENTREE,
        })

        self.assertEqual(len(resultats), 1)
        self.assertEqual(resultats[0].document, "BL-2000")

    def test_valeur_du_mouvement_annotee(self):
        """La valeur est calculée par la base : quantité x prix d'achat."""
        resultats = self._mouvements({"q": "BON-1000"})
        self.assertEqual(resultats[0].valeur, Decimal("5000.00"))  # 5 x 1000

    def test_filtre_utilisateur_invalide_ignore(self):
        resultats = self._mouvements({"utilisateur": "abc"})
        self.assertEqual(len(resultats), 3)

    def test_liste_des_utilisateurs_limitee_aux_actifs(self):
        """Seuls les utilisateurs ayant enregistré un mouvement sont proposés."""
        creer_utilisateur("jamais_actif", groupe="Magasinier")

        reponse = self.client.get(reverse("inventory:mouvement_liste"))
        noms = [u.username for u in reponse.context["utilisateurs"]]

        self.assertIn("gerant", noms)
        self.assertIn("magasinier", noms)
        self.assertNotIn("jamais_actif", noms)


@override_settings(**PARAMETRES_EMAIL_TEST)
class RapportAvecAjustementsTests(BaseApplicationTestCase):
    """Le rapport quotidien compte aussi les ajustements."""

    def setUp(self):
        self.categorie = Categorie.objects.create(nom="Alimentation")
        self.produit = Produit.objects.create(
            reference="RAJ-001", nom="Riz", categorie=self.categorie,
            prix_achat=Decimal("1000.00"), prix_vente=Decimal("1400.00"), quantite_stock=100,
        )

    def test_statistiques_du_jour_incluent_les_ajustements(self):
        services.enregistrer_mouvement(self.produit, Mouvement.ENTREE, 10)
        services.enregistrer_mouvement(self.produit, Mouvement.SORTIE, 4)
        services.enregistrer_mouvement(self.produit, Mouvement.AJUSTEMENT, 2, motif="Casse")

        stats = statistiques.statistiques_du_jour()

        self.assertEqual(stats["ajustements_nombre"], 1)
        self.assertEqual(stats["ajustements_quantite"], 2)
        # Le total additionne bien les trois types.
        self.assertEqual(stats["mouvements_nombre"], 3)

    def test_rapport_email_mentionne_les_ajustements(self):
        services.enregistrer_mouvement(self.produit, Mouvement.AJUSTEMENT, 2, motif="Casse")
        mail.outbox = []

        call_command("rapport_quotidien", verbosity=0)

        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Ajustements", mail.outbox[0].alternatives[0][0])
        self.assertIn("Ajustements", mail.outbox[0].body)
